"""
Document processing service.

Handles page-by-page text extraction from various file formats (PDF, PPTX, DOCX, images),
section/heading hierarchy tracking, and token-based Parent-Child text splitting.
"""

import logging
import json
import re
from pathlib import Path

import pymupdf4llm
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from PIL import Image
from pptx import Presentation

from app.services.vector_store import _sanitize_name

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> list[dict]:
    """
    Extract text from a PDF file page-by-page using pymupdf4llm's JSON output.

    Args:
        file_path: Absolute path to the PDF file.

    Returns:
        List of dictionaries with 'text' and 'page_number' (1-indexed).
    """
    logger.info("Extracting structured page text from PDF: %s", file_path)
    pages = []
    try:
        json_data = pymupdf4llm.to_json(file_path)
        if json_data:
            data = json.loads(json_data)
            for page_data in data.get("pages", []):
                pages.append({
                    "text": page_data.get("fulltext", "").strip(),
                    "page_number": page_data.get("page_number", len(pages) + 1)
                })
    except Exception as e:
        logger.error("Failed to extract pages from PDF via to_json: %s. Falling back to to_markdown.", str(e))
        try:
            text = pymupdf4llm.to_markdown(file_path)
            pages.append({
                "text": text,
                "page_number": 1
            })
        except Exception as fallback_err:
            logger.error("PDF fallback to_markdown failed: %s", str(fallback_err))
    return pages


def extract_text_from_pptx(file_path: str) -> list[dict]:
    """
    Extract text from a PowerPoint file slide-by-slide.

    Args:
        file_path: Absolute path to the PPTX file.

    Returns:
        List of dictionaries with 'text' and 'page_number' (slide number).
    """
    logger.info("Extracting slide text from PPTX: %s", file_path)
    prs = Presentation(file_path)
    pages = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_text_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    paragraph_text = paragraph.text.strip()
                    if paragraph_text:
                        slide_text_parts.append(paragraph_text)

            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        slide_text_parts.append(row_text)

        slide_text = "\n".join(slide_text_parts).strip()
        pages.append({
            "text": slide_text,
            "page_number": slide_num
        })

    return pages


def extract_text_from_docx(file_path: str) -> list[dict]:
    """
    Extract text from a Word document and partition into virtual pages.

    Args:
        file_path: Absolute path to the DOCX file.

    Returns:
        List of dictionaries with 'text' and 'page_number' (virtual page number).
    """
    logger.info("Extracting text from DOCX: %s", file_path)
    doc = DocxDocument(file_path)
    
    pages = []
    current_page_text = []
    current_len = 0
    page_num = 1
    
    for para in doc.paragraphs:
        para_text = para.text.strip()
        if not para_text:
            continue
        current_page_text.append(para_text)
        current_len += len(para_text)
        
        # Partition every ~4000 characters to form virtual pages
        if current_len >= 4000:
            pages.append({
                "text": "\n\n".join(current_page_text),
                "page_number": page_num
            })
            current_page_text = []
            current_len = 0
            page_num += 1
            
    if current_page_text:
        pages.append({
            "text": "\n\n".join(current_page_text),
            "page_number": page_num
        })
        
    return pages


def extract_text_from_txt(file_path: str) -> list[dict]:
    """
    Extract text from a plain text file.

    Args:
        file_path: Absolute path to the text file.

    Returns:
        List of dictionaries with 'text' and 'page_number' (1-indexed).
    """
    logger.info("Extracting text from TXT: %s", file_path)
    text = ""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read().strip()
    except Exception as e:
        logger.error("Failed to extract text from TXT %s: %s", file_path, str(e))
        text = f"[Text extraction failed for {Path(file_path).name}: {str(e)}]"
        
    return [{
        "text": text,
        "page_number": 1
    }]

def extract_text_from_image(file_path: str) -> list[dict]:
    """
    Extract text from an image using OCR. Treats image as a single page.

    Args:
        file_path: Absolute path to the image file.

    Returns:
        List of dictionaries containing OCR text and page number 1.
    """
    logger.info("Extracting text from image: %s", file_path)
    text = ""
    try:
        import pytesseract
        image = Image.open(file_path)
        text = pytesseract.image_to_string(image).strip()
    except Exception as e:
        logger.warning(
            "Tesseract OCR failed for %s: %s. Ensure tesseract is installed on your system.",
            file_path,
            str(e),
        )
        text = f"[OCR extraction failed for {Path(file_path).name}: {str(e)}]"
        
    return [{
        "text": text,
        "page_number": 1
    }]


# Mapping of supported file extensions to their extraction functions
_EXTRACTORS = {
    ".pdf": extract_text_from_pdf,
    ".pptx": extract_text_from_pptx,
    ".docx": extract_text_from_docx,
    ".png": extract_text_from_image,
    ".jpg": extract_text_from_image,
    ".jpeg": extract_text_from_image,
    ".txt": extract_text_from_txt,
}

SUPPORTED_EXTENSIONS = set(_EXTRACTORS.keys())


def process_file(file_path: str) -> list[dict]:
    """
    Process a file and extract its page-by-page text content.

    Args:
        file_path: Absolute path to the file to process.

    Returns:
        List of dictionaries with 'text' and 'page_number'.
    """
    ext = Path(file_path).suffix.lower()

    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(
            f"Unsupported file extension: '{ext}'. "
            f"Supported extensions: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    return extractor(file_path)


def chunk_text(
    pages: list[dict],
    subject: str,
    topic: str,
    source_file: str,
) -> list[Document]:
    """
    Split extracted pages text into Parent-Child chunks and wrap them as LangChain Document objects.
    
    Uses token-based RecursiveCharacterTextSplitter.
    Parent chunk: 1000-1200 tokens
    Child chunk: 200-250 tokens, 20% overlap (50 tokens)
    Each child chunk keeps its parent chunk's text and metadata.

    Args:
        pages: List of dictionaries with 'text' and 'page_number'.
        subject: Subject name for metadata.
        topic: Topic name for metadata.
        source_file: Original filename for metadata.

    Returns:
        List of LangChain Document objects with parent-child structure.
    """
    from config import get_settings
    
    settings = get_settings()
    
    try:
        import tiktoken
        logger.info("Using Tiktoken encoder for Parent-Child text splitting.")
        parent_splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
            model_name=settings.EMBEDDING_MODEL,
            chunk_size=1200,
            chunk_overlap=120,
        )
        child_splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
            model_name=settings.EMBEDDING_MODEL,
            chunk_size=250,
            chunk_overlap=50,
        )
    except ImportError:
        logger.warning("Tokenizer not available, falling back to character-based splitting.")
        parent_splitter = RecursiveCharacterTextSplitter(
            chunk_size=4800,
            chunk_overlap=480,
            separators=["\n\n", "\n", ". ", " ", ""],
        )
        child_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            separators=["\n\n", "\n", ". ", " ", ""],
        )
        
    documents = []
    current_headings = {"h1": "", "h2": "", "h3": ""}
    
    subject_sanit = _sanitize_name(subject)
    topic_sanit = _sanitize_name(topic)
    file_sanit = _sanitize_name(source_file)
    
    for page_data in pages:
        page_text = page_data["text"]
        page_num = page_data["page_number"]
        
        if not page_text.strip():
            continue
            
        # Split page text into parent chunks
        parent_chunks = parent_splitter.split_text(page_text)
        
        for p_idx, parent_chunk_text in enumerate(parent_chunks):
            # Parse headings inside the parent chunk to update heading context
            for line in parent_chunk_text.split("\n"):
                line_strip = line.strip()
                if line_strip.startswith("# "):
                    current_headings["h1"] = line_strip[2:].strip()
                    current_headings["h2"] = ""
                    current_headings["h3"] = ""
                elif line_strip.startswith("## "):
                    current_headings["h2"] = line_strip[3:].strip()
                    current_headings["h3"] = ""
                elif line_strip.startswith("### "):
                    current_headings["h3"] = line_strip[4:].strip()
            
            section_name = current_headings["h3"] or current_headings["h2"] or current_headings["h1"] or "General"
            
            # Split this parent chunk into child chunks
            child_chunks = child_splitter.split_text(parent_chunk_text)
            
            for c_idx, child_chunk_text in enumerate(child_chunks):
                # Update headings if child starts with a heading
                for line in child_chunk_text.split("\n"):
                    line_strip = line.strip()
                    if line_strip.startswith("# "):
                        current_headings["h1"] = line_strip[2:].strip()
                        current_headings["h2"] = ""
                        current_headings["h3"] = ""
                    elif line_strip.startswith("## "):
                        current_headings["h2"] = line_strip[3:].strip()
                        current_headings["h3"] = ""
                    elif line_strip.startswith("### "):
                        current_headings["h3"] = line_strip[4:].strip()
                        
                child_section = current_headings["h3"] or current_headings["h2"] or current_headings["h1"] or section_name
                
                parent_id = f"{subject_sanit}_{topic_sanit}_{file_sanit}_p{page_num}_parent{p_idx}"
                
                doc = Document(
                    page_content=child_chunk_text,
                    metadata={
                        "subject": subject,
                        "topic": topic,
                        "source_file": source_file,
                        "page": page_num,
                        "section": child_section,
                        "parent_id": parent_id,
                        "parent_content": parent_chunk_text,
                        "chunk_index": len(documents),
                    }
                )
                documents.append(doc)
                
    logger.info(
        "Split %s into %d Parent-Child documents.",
        source_file,
        len(documents),
    )
    return documents
