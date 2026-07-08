import requests
import os

BASE_URL = "http://127.0.0.1:8001/api"
SUBJECT = "Science"
TOPIC = "Physics"

def test_teacher_upload():
    print(f"--- 1. Testing Teacher Upload ({SUBJECT} / {TOPIC}) ---")
    
    # Create a dummy Word document to upload
    import docx
    file_path = "dummy_physics_notes.docx"
    doc = docx.Document()
    doc.add_heading('Physics Notes', 0)
    doc.add_paragraph("Physics is the natural science that studies matter, its fundamental constituents, its motion and behavior through space and time, and the related entities of energy and force. Newton's laws of motion are three basic laws of classical mechanics.")
    doc.save(file_path)
        
    url = f"{BASE_URL}/teacher/upload"
    
    # Send multipart/form-data request
    with open(file_path, "rb") as f:
        files = {"files": (file_path, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {"subject": SUBJECT, "topic": TOPIC}
        response = requests.post(url, data=data, files=files)
        
    print(f"Status Code: {response.status_code}")
    print(f"Response: {response.json()}\n")
    
    # Cleanup dummy file
    os.remove(file_path)

def test_list_topics():
    print("--- 2. Testing List Topics ---")
    url = f"{BASE_URL}/student/topics"
    response = requests.get(url)
    print(f"Status Code: {response.status_code}")
    print(f"Response: {response.json()}\n")

def test_generate_notes():
    print(f"--- 3. Testing Generate Notes ({SUBJECT} / {TOPIC}) ---")
    url = f"{BASE_URL}/student/generate"
    
    payload = {
        "subject": SUBJECT,
        "topic": TOPIC,
        "type": "notes",
        "notes_focus": "Default",
        "force_regenerate": True
    }
    
    print("Waiting for Gemini LLM to generate notes (this might take 10-20 seconds)...")
    response = requests.post(url, json=payload)
    
    print(f"Status Code: {response.status_code}")
    if response.status_code == 200:
        content = response.json().get("content", "")
        print(f"Success! Generated {len(content)} characters of notes.")
        print("-" * 50)
        print(content[:500] + "...\n[TRUNCATED FOR DISPLAY]")
        print("-" * 50)
    else:
        print(f"Error: {response.json()}")

if __name__ == "__main__":
    test_teacher_upload()
    test_list_topics()
    test_generate_notes()
